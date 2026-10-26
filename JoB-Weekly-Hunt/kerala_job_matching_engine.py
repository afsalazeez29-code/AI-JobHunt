#!/usr/bin/env python3
"""Deterministic, explainable shortlist generator for existing Kerala job data."""
from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
from copy import copy
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zipfile import BadZipFile
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation


PROFILE_VERSION = "Afsal A Azeez project evidence v1"
WEIGHTS = {"role": 25, "technical": 30, "experience": 15, "project": 15, "education": 10, "practical": 5}
STRONG_THRESHOLD, RELATED_THRESHOLD, CLOSING_SOON_DAYS = 75, 55, 3
REVIEW_OPTIONS = ["Not Reviewed", "Apply", "Review", "Skip", "Applied"]

# Canonical forms are configurable in this one section. Only skills linked to
# supplied resume/project evidence are treated as demonstrated matching skills.
SKILL_ALIASES = {
    "JavaScript": ["javascript", "js"], "TypeScript": ["typescript", "ts"],
    "React": ["react", "react.js", "reactjs"], "Next.js": ["next.js", "nextjs"],
    "Astro": ["astro"], "Svelte": ["svelte", "svelte 5"], "Vite": ["vite"],
    "React Router": ["react router"], "SCSS": ["scss", "sass"], "Tailwind CSS": ["tailwind", "tailwind css"],
    "Figma": ["figma"], "Node.js": ["node.js", "nodejs", "node"], "Express.js": ["express", "express.js"],
    "REST API": ["rest api", "restful"], "JWT": ["jwt", "json web token"], "RBAC": ["rbac", "role-based access"],
    "Zod": ["zod"], "MongoDB": ["mongodb", "mongo"], "Mongoose": ["mongoose"],
    "PostgreSQL": ["postgresql", "postgres"], "Prisma": ["prisma"], "SQL": ["sql"],
    "AWS EC2": ["aws ec2", "ec2"], "Vercel": ["vercel"], "Cloudinary": ["cloudinary"],
    "Cloudflare Pages": ["cloudflare pages"], "Ubuntu Linux": ["ubuntu", "linux"], "PM2": ["pm2"],
    "Nginx": ["nginx"], "Python": ["python"], "Git": ["git"], "Postman": ["postman"],
    "Three.js": ["three.js", "threejs"], "Framer Motion": ["framer motion"], "GSAP": ["gsap"],
}
DEMONSTRATED_SKILLS = set(SKILL_ALIASES)
PROJECT_EVIDENCE = {
    "frontend": "Recipe.IO and JailMeet 2.0 demonstrate responsive React/Next.js web interfaces.",
    "fullstack": "Recipe.IO and JailMeet 2.0 demonstrate Node.js/Express APIs and database-backed workflows.",
    "backend": "JailMeet 2.0 demonstrates protected REST APIs, JWT, RBAC, Zod, PostgreSQL, and Prisma.",
    "deployment": "Recipe.IO demonstrates deployment with AWS EC2, PM2, Nginx, and Vercel.",
}
ROLE_PATTERNS = {
    "Primary — Frontend / Full-Stack": [r"\b(front[ -]?end|react(?:\.js)?|next(?:\.js)?|javascript|typescript|full[ -]?stack|mern|web application|web developer)\b"],
    "Related — Software Development": [r"\b(software (developer|engineer)|application developer|node(?:\.js)? developer|express(?:\.js)? developer|web ui|ui engineer|product engineer|developer trainee|graduate software|developer intern)\b"],
    "Conditional — Transferable Development": [r"\b(python developer|ai application|ui/ux developer|ui ux developer|cloud.*developer|graduate trainee)\b"],
}
EXCLUDED_TITLE_PATTERN = re.compile(r"\b(sales|marketing|recruit|recruitment|hr\b|human resources|finance|accountant|customer support|content writer|legal|business development|graphic designer|animation)\b", re.I)
SENIOR_PATTERN = re.compile(r"\b(senior|lead|principal|architect|manager|head of)\b", re.I)
OUTPUT_COLUMNS = [
    ("Review Status", "review_status"), ("Priority", "priority"), ("Match Score", "match_score"),
    ("Evidence Coverage", "evidence_coverage"), ("Match Classification", "classification"),
    ("Target Role Family", "role_family"), ("Job Title", "job_title"), ("Company", "company"),
    ("Source Portal", "source_portal"), ("Posted Date", "posted_date"), ("Closing Date", "closing_date"),
    ("Deadline Status", "deadline_status"), ("Days Remaining", "days_remaining"), ("Location", "location"),
    ("Experience Requirement", "experience_requirement"), ("Education Requirement", "education_requirement"),
    ("Education Eligibility", "education_eligibility"), ("Experience Eligibility", "experience_eligibility"),
    ("Required Skills", "required_skills"), ("Matching Skills", "matching_skills"),
    ("Missing Information", "missing_information"), ("Mandatory Requirement Conflicts", "mandatory_conflicts"),
    ("Match Explanation", "match_explanation"), ("Matching Method", "matching_method"),
    ("Job Details URL", "job_details_url"), ("Original Job ID", "source_record_id"),
    ("Job Listing ID", "job_listing_id"), ("Notes", "notes"),
]


def now_iso() -> str:
    return datetime.now(ZoneInfo("Asia/Kolkata")).isoformat()


def today_india() -> date:
    return datetime.now(ZoneInfo("Asia/Kolkata")).date()


def text(value: Any) -> str:
    if value is None:
        return ""
    return value if isinstance(value, str) else str(value)


def atomic_json(path: Path, data: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_name(f".{path.name}.tmp")
    temp.write_text(json.dumps(data, ensure_ascii=False, indent=2, default=str) + "\n", encoding="utf-8")
    temp.replace(path)


def source_key(job: dict[str, Any]) -> str:
    return f"{job['source_portal']}:{job['source_record_id']}"


def load_json(path: Path) -> Any:
    return json.loads(path.read_text(encoding="utf-8"))


def discover_source(project_dir: Path, portal: str, explicit: Path | None) -> Path | None:
    if explicit:
        return explicit if explicit.is_file() else None
    name = "technopark_all_jobs.json" if portal == "Technopark" else "infopark_all_jobs.json"
    candidate = project_dir / "data" / portal.lower() / name
    return candidate if candidate.is_file() else None


def normalize_technopark(path: Path) -> list[dict[str, Any]]:
    payload = load_json(path)
    if not isinstance(payload, dict) or not isinstance(payload.get("jobs"), list):
        raise ValueError("Technopark JSON must be an object with a jobs list")
    detail_ok = bool(payload.get("collection_metadata", {}).get("job_details_url_validation", {}).get("validated"))
    jobs = []
    for record in payload["jobs"]:
        if not isinstance(record, dict):
            continue
        company = record.get("company") if isinstance(record.get("company"), dict) else {}
        identifier = record.get("id") if record.get("id") is not None else record.get("job_listing_id")
        if identifier is None:
            identifier = f"missing-{len(jobs) + 1}"
        jobs.append({
            "source_portal": "Technopark", "source_record_id": text(identifier), "job_listing_id": text(record.get("job_listing_id")),
            "job_title": text(record.get("job_title")), "company": text(company.get("company")),
            "posted_date": text(record.get("posted_date")), "closing_date": text(record.get("closing_date")),
            "job_description": "", "required_skills": "", "preferred_skills": "", "experience_requirement": "",
            "education_requirement": "", "employment_type": "", "walk_in_status": text(record.get("is_walk_in")),
            "walk_in_date_time": " ".join(filter(None, [text(record.get("walk_in_start_date")), text(record.get("walk_in_start_time"))])),
            "location": "", "work_mode": "", "salary": "",
            "job_details_url": f"https://technopark.in/job-details/{identifier}" if detail_ok and record.get("id") is not None else "",
            "original_source_file": str(path), "original_source_record": record,
        })
    return jobs


def normalize_infopark(path: Path) -> list[dict[str, Any]]:
    payload = load_json(path)
    if not isinstance(payload, list):
        raise ValueError("Infopark JSON must be a list of job records")
    jobs = []
    for record in payload:
        if not isinstance(record, dict):
            continue
        identifier = record.get("job_id") or record.get("job_details_url") or f"missing-{len(jobs) + 1}"
        jobs.append({
            "source_portal": "Infopark", "source_record_id": text(identifier), "job_listing_id": "",
            "job_title": text(record.get("job_title")), "company": text(record.get("company_name")),
            "posted_date": text(record.get("date_posted")), "closing_date": text(record.get("application_deadline")),
            "job_description": text(record.get("job_description")), "required_skills": "", "preferred_skills": "",
            "experience_requirement": "", "education_requirement": "", "employment_type": "", "walk_in_status": "",
            "walk_in_date_time": "", "location": "", "work_mode": "", "salary": "",
            "job_details_url": text(record.get("job_details_url")), "original_source_file": str(path), "original_source_record": record,
        })
    return jobs


def skill_mentions(value: str) -> set[str]:
    lowered = value.casefold()
    found: set[str] = set()
    for canonical, aliases in SKILL_ALIASES.items():
        if any(re.search(r"(?<![\w.])" + re.escape(alias.casefold()) + r"(?![\w.])", lowered) for alias in aliases):
            found.add(canonical)
    return found


def role_family(title: str) -> tuple[str, int]:
    if EXCLUDED_TITLE_PATTERN.search(title):
        return "Excluded — Unrelated", 0
    for family, patterns in ROLE_PATTERNS.items():
        if any(re.search(pattern, title, re.I) for pattern in patterns):
            return family, 25 if family.startswith("Primary") else 18 if family.startswith("Related") else 14
    return "Excluded — Outside target role taxonomy", 0


def requirement_sections(description: str) -> tuple[set[str], set[str]]:
    """Separate obvious required/preferred phrases without inventing requirements."""
    required, preferred = set(), set()
    for line in re.split(r"[\n.;]", description):
        skills = skill_mentions(line)
        low = line.casefold()
        if re.search(r"\b(preferred|nice to have|bonus|plus)\b", low):
            preferred |= skills
        elif re.search(r"\b(required|must have|requirements?|mandatory)\b", low):
            required |= skills
    return required, preferred


def extract_experience(text_value: str) -> str:
    match = re.search(r"\b(\d+(?:\s*-\s*\d+)?\s*\+?\s*(?:years?|yrs?))\b", text_value, re.I)
    if match:
        return match.group(1)
    if re.search(r"\b(fresher|intern(ship)?|trainee|entry[ -]?level|junior)\b", text_value, re.I):
        return "Entry-level / fresher wording present"
    return ""


def extract_education(text_value: str) -> str:
    match = re.search(r"\b(B\.?(?:Tech|E|Sc|CA)|M\.?(?:Tech|CA|Sc)|Bachelor(?:'s)?|Master(?:'s)?|degree|diploma)\b[^.\n;]*", text_value, re.I)
    return match.group(0).strip() if match else ""


def deadline(closing: str, reference: date) -> tuple[str, int | None]:
    if not closing:
        return "Unknown", None
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d %b %Y", "%d %B %Y"):
        try:
            days = (datetime.strptime(closing.strip(), fmt).date() - reference).days
            return ("Expired" if days < 0 else "Closing Soon" if days <= CLOSING_SOON_DAYS else "Open"), days
        except ValueError:
            pass
    return "Unknown", None


def evaluate(job: dict[str, Any], reference: date) -> dict[str, Any]:
    """Score one normalized job using only title, description, and stated data."""
    title, description = job["job_title"], job["job_description"]
    family, role_score = role_family(title)
    if role_score == 0:
        job.update({"role_family": family, "classification": "Low Match / Excluded", "match_score": 0,
                    "evidence_coverage": "0%", "confidence": "Low", "matching_method": "Title-Based",
                    "match_explanation": "The job title is outside the configured junior web/software target taxonomy.",
                    "missing_information": "Job description and requirements unavailable.", "mandatory_conflicts": "",
                    "experience_eligibility": "Not Specified", "education_eligibility": "Not Specified", "technical_match": "Not Applicable",
                    "overall_eligibility": "Not Specified",
                    "required_skills": "", "preferred_skills": "", "matching_skills": "", "deadline_status": deadline(job["closing_date"], reference)[0],
                    "days_remaining": deadline(job["closing_date"], reference)[1], "priority": "Excluded"})
        return job
    all_text = f"{title}\n{description}"
    mentioned = skill_mentions(all_text)
    required, preferred = requirement_sections(description)
    matching = mentioned & DEMONSTRATED_SKILLS
    required_missing = required - DEMONSTRATED_SKILLS
    exp = extract_experience(description)
    edu = extract_education(description)
    seniority_text = extract_experience(all_text)
    conflicts: list[str] = []
    if re.search(r"\b[2-9]\d*\s*\+?\s*(?:years?|yrs?)\b", seniority_text, re.I):
        conflicts.append(f"Professional experience requirement ({seniority_text}) is not established by the candidate profile.")
    if required_missing:
        conflicts.append("Required skills not in demonstrated profile: " + ", ".join(sorted(required_missing)))
    experience_eligibility = "Appears Eligible" if exp and "Entry-level" in exp else "Needs Verification" if exp else "Not Specified"
    education_eligibility = "Needs Verification" if edu else "Not Specified"
    overall_eligibility = "Needs Verification" if "Needs Verification" in {experience_eligibility, education_eligibility} else "Not Specified"
    technical_score = min(30, len(matching) * 7 + (5 if required and not required_missing else 0))
    project_score = 15 if family.startswith("Primary") else 12 if family.startswith("Related") else 8
    experience_score = 15 if experience_eligibility == "Appears Eligible" else 8 if experience_eligibility == "Not Specified" else 4
    education_score = 10 if not edu else 4
    practical_score = 5 if job["location"] or job["work_mode"] else 3
    score = min(100, role_score + technical_score + project_score + experience_score + education_score + practical_score)
    coverage_units = 1 + bool(description) + bool(required or preferred) + bool(exp) + bool(edu)
    coverage = coverage_units / 5
    method = "Mixed" if description else "Title-Based"
    missing: list[str] = []
    if not description: missing.append("Job description and detailed requirements unavailable in source dataset")
    if not exp: missing.append("Professional experience requirement not specified")
    if not edu: missing.append("Education requirement not specified")
    if not job["location"]: missing.append("Location not specified")
    if SENIOR_PATTERN.search(title): missing.append("Title signals seniority; stated experience requirements need human verification")
    if conflicts:
        classification = "Manual Review" if not description else "Low Match / Excluded"
    elif SENIOR_PATTERN.search(title):
        classification = "Manual Review"
    elif not description:
        classification = "Manual Review"
    elif score >= STRONG_THRESHOLD and coverage >= 0.6:
        classification = "Strong Match"
    elif score >= RELATED_THRESHOLD:
        classification = "Related Match"
    else:
        classification = "Manual Review"
    project_note = PROJECT_EVIDENCE["frontend" if family.startswith("Primary") else "fullstack" if family.startswith("Related") else "backend"]
    evidence = []
    evidence.append(f"Title maps to {family.replace(' — ', ' / ')}.")
    if matching: evidence.append("Mentioned matching skills: " + ", ".join(sorted(matching)) + ".")
    evidence.append(project_note)
    if not description: evidence.append("Assessment is title-based because the source has no job description.")
    status, days = deadline(job["closing_date"], reference)
    priority = "High" if classification == "Strong Match" and status != "Expired" else "Normal" if classification == "Related Match" else "Review" if classification == "Manual Review" else "Excluded"
    job.update({
        "role_family": family, "classification": classification, "match_score": score, "evidence_coverage": f"{round(coverage * 100)}%",
        "confidence": "High" if coverage >= .8 else "Medium" if coverage >= .5 else "Low", "matching_method": method,
        "match_explanation": " ".join(evidence), "missing_information": "; ".join(missing),
        "mandatory_conflicts": "; ".join(conflicts), "experience_eligibility": experience_eligibility,
        "education_eligibility": education_eligibility, "overall_eligibility": overall_eligibility,
        "technical_match": "Matching skills found" if matching else "Title relevance only",
        "required_skills": ", ".join(sorted(required)), "preferred_skills": ", ".join(sorted(preferred)),
        "matching_skills": ", ".join(sorted(matching)), "deadline_status": status, "days_remaining": days, "priority": priority,
    })
    return job


def potential_duplicates(jobs: list[dict[str, Any]]) -> list[dict[str, str]]:
    buckets: dict[tuple[str, str], list[dict[str, Any]]] = {}
    for job in jobs:
        key = (re.sub(r"\W+", "", job["job_title"].casefold()), re.sub(r"\W+", "", job["company"].casefold()))
        if all(key): buckets.setdefault(key, []).append(job)
    return [{"title": group[0]["job_title"], "company": group[0]["company"], "source_keys": ", ".join(source_key(item) for item in group)}
            for group in buckets.values() if len({item["source_portal"] for item in group}) > 1]


def read_review_data(workbook_path: Path, state_path: Path, backup_dir: Path, stamp: str) -> tuple[dict[str, dict[str, str]], list[str]]:
    reviews: dict[str, dict[str, str]] = {}
    warnings: list[str] = []
    if state_path.exists():
        try:
            stored = load_json(state_path)
            if isinstance(stored, dict): reviews.update({k: v for k, v in stored.get("reviews", {}).items() if isinstance(v, dict)})
        except (OSError, ValueError) as exc: warnings.append(f"Review-state file could not be read: {exc}")
    if not workbook_path.exists(): return reviews, warnings
    backup_dir.mkdir(parents=True, exist_ok=True)
    backup = backup_dir / f"kerala_jobs_resume_matched_before_refresh_{stamp}.xlsx"
    shutil.copy2(workbook_path, backup)
    try:
        wb = load_workbook(workbook_path, read_only=True, data_only=False)
        ws = wb["All Shortlisted Jobs"] if "All Shortlisted Jobs" in wb.sheetnames else None
        if not ws: raise ValueError("All Shortlisted Jobs worksheet is missing")
        heads = {text(c.value): i for i, c in enumerate(ws[1])}
        for required in ("Source Portal", "Original Job ID", "Review Status", "Notes"):
            if required not in heads: raise ValueError(f"missing workbook column {required}")
        for row in ws.iter_rows(min_row=2, values_only=True):
            key = f"{text(row[heads['Source Portal']])}:{text(row[heads['Original Job ID']])}"
            if key != ":": reviews[key] = {"review_status": text(row[heads["Review Status"]]), "notes": text(row[heads["Notes"]])}
        wb.close()
    except (OSError, ValueError, KeyError, BadZipFile, InvalidFileException) as exc:
        warnings.append(f"Existing workbook was backed up but could not be read safely: {type(exc).__name__}: {exc}")
    return reviews, warnings


def write_csv(path: Path, jobs: list[dict[str, Any]]) -> None:
    headers = [field for _, field in OUTPUT_COLUMNS]
    with path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=headers)
        writer.writeheader(); writer.writerows([{key: text(job.get(key)) for key in headers} for job in jobs])


def format_sheet(ws, jobs: list[dict[str, Any]]) -> None:
    headers = [head for head, _ in OUTPUT_COLUMNS]; ws.append(headers)
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF"); cell.fill = fill
        align = copy(cell.alignment); align.wrap_text = True; align.vertical = "center"; cell.alignment = align
    for job in jobs:
        row = [job.get(field) for _, field in OUTPUT_COLUMNS]; ws.append(row)
        review_cell = ws.cell(ws.max_row, 1); review_cell.value = job.get("review_status", "Not Reviewed")
        url = job.get("job_details_url")
        if url: ws.cell(ws.max_row, 25).hyperlink = url; ws.cell(ws.max_row, 25).style = "Hyperlink"
    last = max(2, ws.max_row); ws.freeze_panes = "A2"; ws.auto_filter.ref = f"A1:AB{last}"
    validation = DataValidation(type="list", formula1='"' + ",".join(REVIEW_OPTIONS) + '"', allow_blank=False)
    ws.add_data_validation(validation); validation.add(f"A2:A{last}")
    for col, width in {"A":16,"B":11,"C":11,"D":15,"E":18,"F":30,"G":32,"H":30,"I":13,"J":13,"K":13,"L":14,"M":13,"N":18,"O":22,"P":22,"Q":18,"R":20,"S":25,"T":25,"U":36,"V":36,"W":60,"X":16,"Y":46,"Z":16,"AA":20,"AB":40}.items(): ws.column_dimensions[col].width = width
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            align = copy(cell.alignment); align.wrap_text = True; align.vertical = "top"; cell.alignment = align
    ws.conditional_formatting.add(f"E2:E{last}", CellIsRule(operator="equal", formula=['"Strong Match"'], fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(f"L2:L{last}", CellIsRule(operator="equal", formula=['"Expired"'], fill=PatternFill("solid", fgColor="FFC7CE")))


def write_workbook(path: Path, shortlisted: list[dict[str, Any]], report: dict[str, Any]) -> None:
    wb = Workbook(); wb.remove(wb.active)
    by_class = {name: [j for j in shortlisted if j["classification"] == name and j["deadline_status"] != "Expired"] for name in ("Strong Match", "Related Match", "Manual Review")}
    for sheet, jobs in [("Strong Matches", by_class["Strong Match"]), ("Related Matches", by_class["Related Match"]), ("Manual Review", by_class["Manual Review"]),
                        ("Expired - Past Deadline", [j for j in shortlisted if j["deadline_status"] == "Expired"]), ("All Shortlisted Jobs", shortlisted)]:
        ws = wb.create_sheet(sheet); format_sheet(ws, jobs)
    summary = wb.create_sheet("Summary"); summary.append(["Kerala Resume-Matched Job Summary", report["execution_status"]]); summary["A1"].font = Font(bold=True, size=14)
    summary.append(["Run timestamp", report["completed_at"]]); summary.append([])
    for heading, mapping in [("Source Portal", report["counts"]["by_source"]), ("Match Classification", report["counts"]["by_classification"]),
                             ("Target Role Family", report["counts"]["by_role_family"]), ("Eligibility Status", report["counts"]["by_eligibility"]), ("Deadline Status", report["counts"]["by_deadline"]),
                             ("Evidence Coverage", report["counts"]["by_coverage"]), ("Matching Method", report["counts"]["by_method"])]:
        summary.append([heading, "Count"])
        for key, value in sorted(mapping.items()): summary.append([key, value])
        summary.append([])
    summary.column_dimensions["A"].width = 52; summary.column_dimensions["B"].width = 20; summary.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True); temp = path.with_name(f".{path.name}.tmp"); wb.save(temp); wb.close(); temp.replace(path)


def counts(jobs: list[dict[str, Any]]) -> dict[str, dict[str, int]]:
    def group(field: str) -> dict[str, int]:
        result: dict[str, int] = {}
        for job in jobs: result[text(job.get(field)) or "Not Specified"] = result.get(text(job.get(field)) or "Not Specified", 0) + 1
        return result
    return {"by_source": group("source_portal"), "by_classification": group("classification"), "by_role_family": group("role_family"), "by_eligibility": group("overall_eligibility"), "by_deadline": group("deadline_status"), "by_coverage": group("evidence_coverage"), "by_method": group("matching_method")}


def validate_outputs(json_path: Path, csv_path: Path, xlsx_path: Path, expected: int) -> dict[str, Any]:
    result = {"passed": True, "checks": [], "errors": []}
    try:
        assert len(load_json(json_path)["shortlisted_jobs"]) == expected; result["checks"].append("JSON reopened with expected shortlist count.")
        with csv_path.open(encoding="utf-8-sig", newline="") as h: rows = list(csv.DictReader(h))
        assert len(rows) == expected; result["checks"].append("CSV reopened with expected shortlist count.")
        wb = load_workbook(xlsx_path); assert {"Strong Matches","Related Matches","Manual Review","Expired - Past Deadline","All Shortlisted Jobs","Summary"}.issubset(wb.sheetnames)
        ws = wb["All Shortlisted Jobs"]; assert ws.max_row - 1 == expected and ws.freeze_panes == "A2" and ws.auto_filter.ref and ws.data_validations.dataValidation
        wb.close(); result["checks"].append("Workbook reopened with required sheets, rows, filters, freeze pane, and review validation.")
    except (OSError, ValueError, AssertionError, KeyError, BadZipFile, InvalidFileException) as exc: result["passed"] = False; result["errors"].append(str(exc))
    return result


def parse_args() -> argparse.Namespace:
    p = argparse.ArgumentParser(description="Create an explainable resume-matched Kerala developer-job shortlist.")
    p.add_argument("--technopark-source", type=Path, help="Explicit Technopark JSON source path.")
    p.add_argument("--infopark-source", type=Path, help="Explicit Infopark JSON source path.")
    p.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent / "data" / "job_matching")
    return p.parse_args()


def main() -> int:
    args = parse_args(); project = Path(__file__).resolve().parent; output = args.output_dir.resolve(); started = now_iso(); warnings: list[str] = []; errors: list[str] = []
    paths = {"Technopark": discover_source(project, "Technopark", args.technopark_source), "Infopark": discover_source(project, "Infopark", args.infopark_source)}
    loaded: list[dict[str, Any]] = []; source_counts = {"Technopark": 0, "Infopark": 0}
    for portal, loader in (("Technopark", normalize_technopark), ("Infopark", normalize_infopark)):
        if not paths[portal]: warnings.append(f"{portal} source dataset is unavailable."); continue
        try: records = loader(paths[portal]); loaded.extend(records); source_counts[portal] = len(records)
        except (OSError, ValueError, json.JSONDecodeError) as exc: errors.append(f"{portal} source could not be loaded: {exc}")
    if not loaded:
        report = {"started_at": started, "completed_at": now_iso(), "execution_status": "FAILED", "warnings": warnings, "errors": errors or ["No valid source datasets loaded."]}
        atomic_json(output / "reports" / "job_matching_report.json", report); return 2
    evaluated = [evaluate(job, today_india()) for job in loaded]
    shortlisted = [job for job in evaluated if job["classification"] != "Low Match / Excluded"]
    duplicates = potential_duplicates(loaded)
    stamp = datetime.now(ZoneInfo("Asia/Kolkata")).strftime("%Y%m%dT%H%M%S%z")
    workbook_path, state_path = output / "kerala_jobs_resume_matched.xlsx", output / "review_state.json"
    reviews, review_warnings = read_review_data(workbook_path, state_path, output / "review_backups", stamp); warnings.extend(review_warnings)
    for job in shortlisted:
        review = reviews.get(source_key(job), {}); job["review_status"] = review.get("review_status", "Not Reviewed"); job["notes"] = review.get("notes", "")
        reviews[source_key(job)] = {"review_status": job["review_status"], "notes": job["notes"]}
    report_counts = counts(evaluated); short_counts = counts(shortlisted)
    status = "COMPLETE" if not warnings and not errors and all(paths.values()) else "PARTIAL"
    report = {"started_at": started, "completed_at": now_iso(), "execution_status": status, "candidate_profile_version": PROFILE_VERSION,
              "matching_configuration": {"weights": WEIGHTS,"strong_threshold":STRONG_THRESHOLD,"related_threshold":RELATED_THRESHOLD,"closing_soon_days":CLOSING_SOON_DAYS,"reference_date":today_india().isoformat()},
              "source_files": {portal: str(path) if path else None for portal,path in paths.items()}, "records_loaded": source_counts,
              "total_source_records_processed": len(loaded), "unique_source_records_processed": len({source_key(j) for j in loaded}),
              "records_missing_stable_identifiers": sum(j["source_record_id"].startswith("missing-") for j in loaded), "potential_cross_source_duplicates": duplicates,
              "strong_match_count": sum(j["classification"] == "Strong Match" for j in shortlisted), "related_match_count": sum(j["classification"] == "Related Match" for j in shortlisted),
              "manual_review_count": sum(j["classification"] == "Manual Review" for j in shortlisted), "expired_matched_job_count": sum(j["deadline_status"] == "Expired" for j in shortlisted),
              "low_match_excluded_count": sum(j["classification"] == "Low Match / Excluded" for j in evaluated),
              "eligibility_uncertainty_count": sum(j["overall_eligibility"] in {"Needs Verification", "Not Specified"} for j in shortlisted),
              "records_without_descriptions": sum(not j["job_description"] for j in evaluated), "title_based_assessments": sum(j["matching_method"] == "Title-Based" for j in shortlisted),
              "counts": short_counts, "all_processed_counts": report_counts, "warnings": warnings, "errors": errors,
              "output_files": {"excel":str(workbook_path),"csv":str(output / "kerala_jobs_resume_matched.csv"),"json":str(output / "kerala_jobs_resume_matched.json"),"review_state":str(state_path),"report":str(output / "reports" / "job_matching_report.json")}}
    json_path, csv_path, report_path = output / "kerala_jobs_resume_matched.json", output / "kerala_jobs_resume_matched.csv", output / "reports" / "job_matching_report.json"
    atomic_json(json_path, {"matching_run": {k: report[k] for k in ("started_at","completed_at","candidate_profile_version","matching_configuration","source_files")}, "shortlisted_jobs": shortlisted})
    output.mkdir(parents=True, exist_ok=True); write_csv(csv_path, shortlisted); write_workbook(workbook_path, shortlisted, report); atomic_json(state_path, {"updated_at":now_iso(), "reviews":reviews})
    report["validation"] = validate_outputs(json_path, csv_path, workbook_path, len(shortlisted))
    if not report["validation"]["passed"]: report["execution_status"] = "FAILED"
    atomic_json(report_path, report)
    print(f"[INFO] Technopark records loaded: {source_counts['Technopark']}"); print(f"[INFO] Infopark records loaded: {source_counts['Infopark']}")
    print(f"[INFO] Shortlisted jobs: {len(shortlisted)}; excluded: {report['low_match_excluded_count']}"); print(f"[INFO] Matching status: {report['execution_status']}")
    return 0 if report["execution_status"] == "COMPLETE" else 2


if __name__ == "__main__": raise SystemExit(main())
