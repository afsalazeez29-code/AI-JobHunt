"""Collect every publicly accessible vacancy from Infopark's All Parks listing.

The collector deliberately does not filter jobs.  A run is only marked complete
when its terminal listing page was verified and no expected listing page failed.
"""
from __future__ import annotations

import argparse
import csv
import json
import logging
import re
import shutil
import sys
import time
from dataclasses import dataclass
from datetime import datetime, timezone
from email.utils import parsedate_to_datetime
from pathlib import Path
from typing import Any, Iterable
from urllib.parse import parse_qs, urlencode, urljoin, urlsplit, urlunsplit

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

SOURCE_URL = "https://infopark.in/companies-job"
SOURCE_HOST = "infopark.in"
FIELDS = [
    "job_id", "job_title", "company_name", "date_posted", "application_deadline",
    "job_details_url", "source_portal", "source_page", "collected_at",
    "collection_run_id", "my_decision", "notes", "validation_issues",
    "job_description", "external_application_url", "detail_extraction_error",
]
REQUIRED_FIELDS = ("job_title", "company_name", "job_details_url")
EXPECTED_HEADERS = ("date of posting", "job title", "company name", "last date to apply")


@dataclass
class ParsedPage:
    records: list[dict[str, str]]
    max_page: int | None
    has_next: bool
    warnings: list[str]


class InfoparkCollector:
    def __init__(self, root: Path, delay: float = 1.0, retries: int = 3, timeout: int = 30,
                 max_pages: int = 250) -> None:
        self.root = root
        self.data_dir = root / "data" / "infopark"
        self.raw_dir = self.data_dir / "raw"
        self.report_dir = self.data_dir / "reports"
        self.log_dir = self.data_dir / "logs"
        for directory in (self.raw_dir, self.report_dir, self.log_dir):
            directory.mkdir(parents=True, exist_ok=True)
        self.delay, self.retries, self.timeout, self.max_pages = delay, retries, timeout, max_pages
        self.session = requests.Session()
        self.session.headers.update({
            "User-Agent": "AI-JobHunt Infopark collector/1.0 (+public listing collection)",
            "Accept": "text/html,application/xhtml+xml",
        })
        self.last_request_at = 0.0
        self.logger = logging.getLogger(f"infopark.{id(self)}")
        self.logger.setLevel(logging.INFO)
        self.logger.propagate = False
        handler = logging.FileHandler(self.log_dir / "infopark_collector.log", encoding="utf-8")
        handler.setFormatter(logging.Formatter("%(asctime)s %(levelname)s %(message)s"))
        self.logger.addHandler(handler)

    @staticmethod
    def timestamp() -> str:
        return datetime.now().astimezone().isoformat(timespec="seconds")

    @staticmethod
    def stamp_for_filename() -> str:
        return datetime.now().astimezone().strftime("%Y-%m-%d_%H%M%S")

    @staticmethod
    def page_url(page: int) -> str:
        return SOURCE_URL if page == 1 else f"{SOURCE_URL}?{urlencode({'page': page})}"

    def _wait_before_request(self) -> None:
        remaining = self.delay - (time.monotonic() - self.last_request_at)
        if remaining > 0:
            time.sleep(remaining)

    def get_text(self, url: str) -> str:
        """Fetch one public portal page with bounded retries and 429 support."""
        last_error: Exception | None = None
        for attempt in range(1, self.retries + 1):
            self._wait_before_request()
            try:
                response = self.session.get(url, timeout=self.timeout)
                self.last_request_at = time.monotonic()
                if response.status_code == 429:
                    retry_after = response.headers.get("Retry-After")
                    wait = self._retry_after_seconds(retry_after) or min(60, 2 ** attempt)
                    self.logger.warning("429 for %s; respecting Retry-After for %.1fs", url, wait)
                    if attempt < self.retries:
                        time.sleep(wait)
                        continue
                if response.status_code in {500, 502, 503, 504} and attempt < self.retries:
                    wait = min(30, 2 ** attempt)
                    self.logger.warning("HTTP %s for %s; retrying in %ss", response.status_code, url, wait)
                    time.sleep(wait)
                    continue
                response.raise_for_status()
                return response.text
            except requests.RequestException as exc:
                last_error = exc
                self.logger.warning("Request %s/%s failed for %s: %s", attempt, self.retries, url, exc)
                if attempt < self.retries:
                    time.sleep(min(30, 2 ** attempt))
        raise RuntimeError(f"Request failed after {self.retries} attempts: {last_error}")

    @staticmethod
    def _retry_after_seconds(value: str | None) -> float | None:
        if not value:
            return None
        try:
            return max(0.0, float(value))
        except ValueError:
            try:
                return max(0.0, (parsedate_to_datetime(value) - datetime.now(timezone.utc)).total_seconds())
            except (TypeError, ValueError):
                return None

    @staticmethod
    def canonical_url(href: str | None, base_url: str = SOURCE_URL) -> str:
        if not href:
            return ""
        absolute = urljoin(base_url, href.strip())
        parts = urlsplit(absolute)
        if parts.scheme not in {"http", "https"}:
            return ""
        # Fragments never identify a distinct portal vacancy; retain query strings conservatively.
        return urlunsplit((parts.scheme.lower(), parts.netloc.lower(), parts.path, parts.query, ""))

    @staticmethod
    def job_id_from_url(url: str) -> str:
        match = re.search(r"/company-jobs/details/\d+/(\d+)(?:/)?$", urlsplit(url).path)
        return match.group(1) if match else ""

    @staticmethod
    def page_number_from_url(url: str) -> int | None:
        try:
            values = parse_qs(urlsplit(url).query).get("page", [])
            number = int(values[-1]) if values else None
            return number if number and number > 0 else None
        except (TypeError, ValueError):
            return None

    @classmethod
    def parse_listing_page(cls, html: str, page_number: int, collected_at: str,
                           run_id: str) -> ParsedPage:
        soup = BeautifulSoup(html, "html.parser")
        table = soup.select_one("#job-list table")
        if not table:
            raise ValueError("Expected #job-list table was not found")
        headers = [cell.get_text(" ", strip=True).casefold() for cell in table.select("thead th")]
        if not all(header in headers for header in EXPECTED_HEADERS):
            raise ValueError(f"Unexpected listing headers: {headers}")
        indices = {header: headers.index(header) for header in EXPECTED_HEADERS}
        records: list[dict[str, str]] = []
        warnings: list[str] = []
        for row_index, row in enumerate(table.select("tbody tr"), start=1):
            cells = row.select("td")
            if not cells or not any(cell.get_text(" ", strip=True) for cell in cells):
                warnings.append(f"page {page_number}, row {row_index}: empty row")
                continue
            if len(cells) <= max(indices.values()):
                warnings.append(f"page {page_number}, row {row_index}: malformed row with {len(cells)} cells")
                continue
            details_link = next((a.get("href") for a in row.select("a[href]")
                                 if "detail" in a.get("href", "").casefold()), None)
            details_url = cls.canonical_url(details_link)
            record = {
                "job_id": cls.job_id_from_url(details_url),
                "job_title": cells[indices["job title"]].get_text(" ", strip=True),
                "company_name": cells[indices["company name"]].get_text(" ", strip=True),
                "date_posted": cells[indices["date of posting"]].get_text(" ", strip=True),
                "application_deadline": cells[indices["last date to apply"]].get_text(" ", strip=True),
                "job_details_url": details_url,
                "source_portal": "Infopark",
                "source_page": str(page_number),
                "collected_at": collected_at,
                "collection_run_id": run_id,
                "my_decision": "",
                "notes": "",
                "validation_issues": "",
                "job_description": "",
                "external_application_url": "",
                "detail_extraction_error": "",
            }
            issues = cls.validate_record(record)
            record["validation_issues"] = "; ".join(issues)
            records.append(record)

        pagination = soup.select_one(".pagination")
        page_numbers: list[int] = []
        has_next = False
        if pagination:
            for node in pagination.select("a[href], span"):
                href = node.get("href", "") if node.name == "a" else ""
                number = cls.page_number_from_url(href) if href else None
                if number:
                    page_numbers.append(number)
                label = " ".join((node.get("rel", []) if isinstance(node.get("rel"), list) else [node.get("rel", "")]))
                label += " " + node.get("aria-label", "") + " " + node.get_text(" ", strip=True)
                if node.name == "a" and ("next" in label.casefold() or node.get("rel") == ["next"]):
                    has_next = True
        return ParsedPage(records, max(page_numbers) if page_numbers else None, has_next, warnings)

    @staticmethod
    def validate_record(record: dict[str, str]) -> list[str]:
        issues = [f"missing {field}" for field in REQUIRED_FIELDS if not record.get(field, "").strip()]
        url = record.get("job_details_url", "")
        parts = urlsplit(url) if url else None
        if url and (not parts or parts.hostname not in {SOURCE_HOST, f"www.{SOURCE_HOST}"}
                    or not parts.path.startswith("/company-jobs/details/")):
            issues.append("detail URL is not an expected official Infopark vacancy URL")
        return issues

    @staticmethod
    def deduplicate(records: Iterable[dict[str, str]]) -> tuple[list[dict[str, str]], int]:
        unique: list[dict[str, str]] = []
        seen: set[str] = set()
        duplicates = 0
        for record in records:
            url = record.get("job_details_url", "")
            # Do not merge malformed rows without a canonical URL.
            key = f"url:{url}" if url else f"unkeyed:{len(unique)}"
            if key in seen:
                duplicates += 1
                continue
            seen.add(key)
            unique.append(record)
        return unique, duplicates

    def discover_pages(self) -> dict[str, Any]:
        started = self.timestamp()
        html = self.get_text(SOURCE_URL)
        parsed = self.parse_listing_page(html, 1, started, "discovery")
        terminal_verified = False
        terminal_warning = ""
        if parsed.max_page and parsed.max_page > 1:
            try:
                final = self.parse_listing_page(self.get_text(self.page_url(parsed.max_page)), parsed.max_page,
                                                started, "discovery")
                terminal_verified = not final.has_next
                if not terminal_verified:
                    terminal_warning = "Highest displayed page still has a next link; sequential fallback is required."
            except Exception as exc:  # discovery must report rather than hide final-page trouble
                terminal_warning = f"Could not verify highest displayed page: {exc}"
        elif parsed.max_page == 1:
            terminal_verified = not parsed.has_next
        return {
            "source_url": SOURCE_URL, "discovered_page_count": parsed.max_page,
            "first_page_has_next": parsed.has_next, "terminal_page_verified": terminal_verified,
            "warnings": parsed.warnings + ([terminal_warning] if terminal_warning else []),
            "checked_at": self.timestamp(),
        }

    def collect_all_pages(self, run_id: str, include_details: bool = False) -> dict[str, Any]:
        started = self.timestamp()
        page_results: dict[int, ParsedPage] = {}
        failed_pages: list[dict[str, Any]] = []
        all_records: list[dict[str, str]] = []
        warnings: list[str] = []
        unexpected_empty_pages: list[int] = []
        try:
            first = self.parse_listing_page(self.get_text(SOURCE_URL), 1, started, run_id)
        except Exception as exc:
            return self._failed_initial_report(started, exc)
        page_results[1] = first
        all_records.extend(first.records)
        warnings.extend(first.warnings)
        expected_pages = first.max_page
        queue = list(range(2, expected_pages + 1)) if expected_pages else ([2] if first.has_next else [])
        attempted = {1}
        while queue:
            page = queue.pop(0)
            if page in attempted:
                continue
            if len(attempted) >= self.max_pages:
                warnings.append(f"Safety limit of {self.max_pages} pages reached; pagination is not verified complete.")
                break
            attempted.add(page)
            url = self.page_url(page)
            try:
                parsed = self.parse_listing_page(self.get_text(url), page, started, run_id)
                page_results[page] = parsed
                all_records.extend(parsed.records)
                warnings.extend(parsed.warnings)
                if not parsed.records:
                    warnings.append(f"page {page}: listing table had no vacancy rows")
                    unexpected_empty_pages.append(page)
                if parsed.max_page and (expected_pages is None or parsed.max_page > expected_pages):
                    expected_pages = parsed.max_page
                    queue.extend(p for p in range(page + 1, expected_pages + 1) if p not in attempted)
                if parsed.has_next and (expected_pages is None or page >= expected_pages):
                    expected_pages = max(expected_pages or 0, page + 1)
                    queue.append(page + 1)
            except Exception as exc:
                failed_pages.append({"page": page, "url": url, "error": str(exc)})
        unique, duplicate_count = self.deduplicate(all_records)
        if include_details:
            self.enrich_records(unique, warnings)
        successful_pages = sorted(page_results)
        terminal_page = expected_pages or max(successful_pages, default=0)
        terminal = page_results.get(terminal_page)
        terminal_verified = bool(terminal and not terminal.has_next)
        expected_attempted = expected_pages is not None and all(page in attempted for page in range(1, expected_pages + 1))
        complete = bool(expected_attempted and terminal_verified and not failed_pages and not unexpected_empty_pages
                        and len(attempted) < self.max_pages)
        return {
            "collection_started_at": started, "collection_finished_at": self.timestamp(), "source_url": SOURCE_URL,
            "pagination": {"pages_expected": expected_pages, "terminal_page": terminal_page,
                           "terminal_page_verified": terminal_verified, "safety_limit": self.max_pages},
            "pages_attempted": sorted(attempted), "pages_successful": successful_pages,
            "failed_pages": failed_pages, "raw_row_count": len(all_records), "unique_vacancy_count": len(unique),
            "duplicate_count": duplicate_count, "malformed_row_count": sum("malformed row" in x for x in warnings),
            "records_missing_required_fields": sum(bool(self.validate_record(r)) for r in unique),
            "unexpected_empty_pages": unexpected_empty_pages,
            "validation_warnings": warnings, "pagination_completion_verified": terminal_verified,
            "complete_collection_success": complete, "records": unique,
        }

    def _failed_initial_report(self, started: str, error: Exception) -> dict[str, Any]:
        return {"collection_started_at": started, "collection_finished_at": self.timestamp(), "source_url": SOURCE_URL,
                "pagination": {"pages_expected": None, "terminal_page": None, "terminal_page_verified": False,
                               "safety_limit": self.max_pages}, "pages_attempted": [1], "pages_successful": [],
                "failed_pages": [{"page": 1, "url": SOURCE_URL, "error": str(error)}], "raw_row_count": 0,
                "unique_vacancy_count": 0, "duplicate_count": 0, "malformed_row_count": 0,
                "records_missing_required_fields": 0, "unexpected_empty_pages": [],
                "validation_warnings": ["Initial listing page was unavailable."],
                "pagination_completion_verified": False, "complete_collection_success": False, "records": []}

    def enrich_records(self, records: list[dict[str, str]], warnings: list[str]) -> None:
        for record in records:
            url = record.get("job_details_url", "")
            if not url:
                record["detail_extraction_error"] = "No official detail URL"
                continue
            try:
                soup = BeautifulSoup(self.get_text(url), "html.parser")
                box = soup.select_one(".comp-job-deatiil .deatil-box, .comp-job-detail .detail-box")
                if not box:
                    raise ValueError("expected detail content container was not found")
                contact = box.select_one(".contact")
                if contact:
                    contact.extract()
                record["job_description"] = box.get_text("\n", strip=True)
                external = next((self.canonical_url(a.get("href"), url) for a in box.select("a[href]")
                                 if urlsplit(urljoin(url, a.get("href", ""))).hostname not in {SOURCE_HOST, f"www.{SOURCE_HOST}"}), "")
                record["external_application_url"] = external
            except Exception as exc:
                record["detail_extraction_error"] = str(exc)
                warnings.append(f"detail extraction failed for {url}: {exc}")

    @staticmethod
    def _record_key(record: dict[str, str]) -> str:
        return f"id:{record['job_id']}" if record.get("job_id") else f"url:{record.get('job_details_url', '')}"

    def existing_manual_values(self) -> dict[str, tuple[str, str]]:
        values: dict[str, tuple[str, str]] = {}
        csv_path = self.data_dir / "infopark_all_jobs.csv"
        if csv_path.exists():
            with csv_path.open("r", encoding="utf-8-sig", newline="") as handle:
                for row in csv.DictReader(handle):
                    values[self._record_key(row)] = (row.get("my_decision", ""), row.get("notes", ""))
        workbook_path = self.data_dir / "infopark_jobs_review.xlsx"
        if workbook_path.exists():
            workbook = load_workbook(workbook_path, read_only=True, data_only=False)
            if "All Jobs" in workbook.sheetnames:
                sheet = workbook["All Jobs"]
                headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
                for row_values in sheet.iter_rows(min_row=2, values_only=True):
                    row = {str(headers[i]): "" if value is None else str(value) for i, value in enumerate(row_values)
                           if headers[i] is not None}
                    key = self._record_key(row)
                    if key != "url:":
                        decision, notes = row.get("my_decision", ""), row.get("notes", "")
                        if decision or notes:
                            values[key] = (decision, notes)
            workbook.close()
        return values

    def existing_primary_records(self) -> list[dict[str, str]]:
        """Read the prior verified CSV solely to report portal disappearances safely."""
        path = self.data_dir / "infopark_all_jobs.csv"
        if not path.exists():
            return []
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            return list(csv.DictReader(handle))

    def apply_manual_values(self, records: list[dict[str, str]]) -> None:
        existing = self.existing_manual_values()
        for record in records:
            decision, notes = existing.get(self._record_key(record), ("", ""))
            record["my_decision"], record["notes"] = decision, notes

    def write_outputs(self, report: dict[str, Any]) -> dict[str, str]:
        """Always preserve a timestamped raw snapshot; only replace primary files if complete."""
        stamp = self.stamp_for_filename()
        records = report.pop("records")
        old_records = self.existing_primary_records()
        self.apply_manual_values(records)
        partial = not report["complete_collection_success"]
        suffix = "_partial" if partial else ""
        snapshot_csv = self.raw_dir / f"infopark_all_jobs_{stamp}{suffix}.csv"
        snapshot_json = self.raw_dir / f"infopark_all_jobs_{stamp}{suffix}.json"
        self._write_csv(snapshot_csv, records)
        self._write_json(snapshot_json, records)
        report["snapshot_csv"] = str(snapshot_csv)
        report["snapshot_json"] = str(snapshot_json)
        report["primary_outputs_updated"] = not partial
        outputs = {"snapshot_csv": str(snapshot_csv), "snapshot_json": str(snapshot_json)}
        if not partial:
            current_keys = {self._record_key(record) for record in records}
            report["listings_disappeared_from_current_portal"] = [
                {"job_id": record.get("job_id", ""), "job_details_url": record.get("job_details_url", ""),
                 "job_title": record.get("job_title", ""), "company_name": record.get("company_name", "")}
                for record in old_records if self._record_key(record) not in current_keys
            ]
            primary_csv = self.data_dir / "infopark_all_jobs.csv"
            primary_json = self.data_dir / "infopark_all_jobs.json"
            review_book = self.data_dir / "infopark_jobs_review.xlsx"
            if review_book.exists():
                backup = self.raw_dir / f"infopark_jobs_review_before_refresh_{stamp}.xlsx"
                shutil.copy2(review_book, backup)
                outputs["workbook_backup"] = str(backup)
            self._write_csv(primary_csv, records)
            self._write_json(primary_json, records)
            self._write_workbook(review_book, records)
            outputs.update({"csv": str(primary_csv), "json": str(primary_json), "xlsx": str(review_book)})
        else:
            report["listings_disappeared_from_current_portal"] = []
        failure_file = self.report_dir / "infopark_failed_pages.json"
        failure_payload = {"failed_pages": report["failed_pages"], "partial_dataset_path": str(snapshot_json),
                           "report_generated_at": self.timestamp(), "previous_report": {
                               "pages_expected": report["pagination"]["pages_expected"],
                               "pages_successful": report["pages_successful"],
                               "pagination_completion_verified": report["pagination_completion_verified"]}}
        self._write_json(failure_file, failure_payload)
        report_path = self.report_dir / "infopark_collection_report.json"
        self._write_json(report_path, report)
        outputs.update({"report": str(report_path), "failures": str(failure_file)})
        return outputs

    @staticmethod
    def _write_csv(path: Path, records: list[dict[str, str]]) -> None:
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=FIELDS, extrasaction="ignore")
            writer.writeheader()
            writer.writerows(records)

    @staticmethod
    def _write_json(path: Path, value: Any) -> None:
        with path.open("w", encoding="utf-8", newline="\n") as handle:
            json.dump(value, handle, ensure_ascii=False, indent=2)

    @staticmethod
    def _write_workbook(path: Path, records: list[dict[str, str]]) -> None:
        workbook = Workbook()
        instructions = workbook.active
        instructions.title = "Instructions"
        instructions.append(["Infopark Job Review Workbook"])
        instructions["A1"].font = Font(bold=True, size=14)
        instructions.append(["Use the All Jobs sheet to search, sort, and filter the complete unfiltered portal dataset."])
        instructions.append(["Set my_decision to Apply, Review, or Skip; leave it blank when undecided. Notes are yours to edit."])
        instructions.append(["Future complete collections preserve decisions and notes by official job ID or detail URL."])
        instructions.column_dimensions["A"].width = 115
        sheet = workbook.create_sheet("All Jobs")
        sheet.append(FIELDS)
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        for record in records:
            sheet.append([record.get(field, "") for field in FIELDS])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:{chr(64 + min(len(FIELDS), 26))}{max(1, sheet.max_row)}"
        widths = {"A": 12, "B": 35, "C": 32, "D": 16, "E": 20, "F": 55, "G": 14, "H": 12,
                  "I": 28, "J": 23, "K": 16, "L": 36, "M": 35, "N": 60, "O": 50, "P": 35}
        for column, width in widths.items():
            sheet.column_dimensions[column].width = width
        decision_col = FIELDS.index("my_decision") + 1
        letter = chr(64 + decision_col)
        validation = DataValidation(type="list", formula1='"Apply,Review,Skip"', allow_blank=True)
        sheet.add_data_validation(validation)
        validation.add(f"{letter}2:{letter}{max(2, sheet.max_row)}")
        sheet.conditional_formatting.add(f"{letter}2:{letter}{max(2, sheet.max_row)}",
                                         FormulaRule(formula=[f'${letter}2="Apply"'], fill=PatternFill("solid", fgColor="C6EFCE")))
        sheet.row_dimensions[1].height = 28
        workbook.save(path)

    def retry_failed_pages(self, run_id: str, include_details: bool = False) -> dict[str, Any]:
        failure_path = self.report_dir / "infopark_failed_pages.json"
        if not failure_path.exists():
            raise FileNotFoundError("No saved failed-page report exists; run a standard collection first.")
        payload = json.loads(failure_path.read_text(encoding="utf-8"))
        old_failures = payload.get("failed_pages", [])
        partial_path = Path(payload.get("partial_dataset_path", ""))
        if not old_failures:
            raise RuntimeError("The most recent collection has no failed pages to retry.")
        if not partial_path.exists():
            raise FileNotFoundError(f"Saved partial dataset is unavailable: {partial_path}")
        started = self.timestamp()
        records = json.loads(partial_path.read_text(encoding="utf-8"))
        still_failed: list[dict[str, Any]] = []
        successes = set(payload.get("previous_report", {}).get("pages_successful", []))
        warnings: list[str] = []
        for failed in old_failures:
            page, url = int(failed["page"]), str(failed["url"])
            try:
                parsed = self.parse_listing_page(self.get_text(url), page, started, run_id)
                records.extend(parsed.records)
                warnings.extend(parsed.warnings)
                successes.add(page)
            except Exception as exc:
                still_failed.append({"page": page, "url": url, "error": str(exc)})
        unique, duplicates = self.deduplicate(records)
        if include_details:
            self.enrich_records(unique, warnings)
        expected = payload.get("previous_report", {}).get("pages_expected")
        terminal_verified = bool(payload.get("previous_report", {}).get("pagination_completion_verified"))
        complete = bool(expected and terminal_verified and not still_failed and
                        set(range(1, int(expected) + 1)).issubset(successes))
        return {"collection_started_at": started, "collection_finished_at": self.timestamp(), "source_url": SOURCE_URL,
                "mode": "retry-failed", "pagination": {"pages_expected": expected, "terminal_page": expected,
                "terminal_page_verified": terminal_verified, "safety_limit": self.max_pages},
                "pages_attempted": sorted({int(x["page"]) for x in old_failures}), "pages_successful": sorted(successes),
                "failed_pages": still_failed, "raw_row_count": len(records), "unique_vacancy_count": len(unique),
                "duplicate_count": duplicates, "malformed_row_count": sum("malformed row" in x for x in warnings),
                "records_missing_required_fields": sum(bool(self.validate_record(r)) for r in unique),
                "unexpected_empty_pages": [],
                "validation_warnings": warnings, "pagination_completion_verified": terminal_verified,
                "complete_collection_success": complete, "records": unique}


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Collect all public Infopark job listings without filtering.")
    group = parser.add_mutually_exclusive_group()
    group.add_argument("--discover-pages", action="store_true", help="Inspect pagination only (first and terminal listing pages).")
    group.add_argument("--retry-failed", action="store_true", help="Retry only pages listed in the last failure report.")
    parser.add_argument("--include-details", action="store_true", help="Also request each official job detail page.")
    parser.add_argument("--delay", type=float, default=1.0, help="Minimum seconds between HTTP requests (default: 1).")
    parser.add_argument("--retries", type=int, default=3, help="Attempts per request (default: 3).")
    parser.add_argument("--timeout", type=int, default=30, help="HTTP timeout in seconds (default: 30).")
    parser.add_argument("--max-pages", type=int, default=250, help="Sequential fallback safety limit (default: 250).")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    if args.delay < 0 or args.retries < 1 or args.max_pages < 1:
        print("delay must be >= 0; retries and max-pages must be >= 1", file=sys.stderr)
        return 2
    collector = InfoparkCollector(Path(__file__).resolve().parent, args.delay, args.retries, args.timeout, args.max_pages)
    try:
        if args.discover_pages:
            print(json.dumps(collector.discover_pages(), ensure_ascii=False, indent=2))
            return 0
        run_id = collector.stamp_for_filename()
        report = collector.retry_failed_pages(run_id, args.include_details) if args.retry_failed else collector.collect_all_pages(run_id, args.include_details)
        outputs = collector.write_outputs(report)
        print(json.dumps({"complete_collection_success": report["complete_collection_success"],
                          "pages_successful": len(report["pages_successful"]), "pages_failed": len(report["failed_pages"]),
                          "raw_rows": report["raw_row_count"], "unique_vacancies": report["unique_vacancy_count"],
                          "duplicates_removed": report["duplicate_count"], "outputs": outputs}, ensure_ascii=False, indent=2))
        return 0 if report["complete_collection_success"] else 1
    except Exception as exc:
        collector.logger.exception("Collector stopped: %s", exc)
        print(f"Collector stopped: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
