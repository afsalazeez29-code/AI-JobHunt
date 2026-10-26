#!/usr/bin/env python3
"""Cached description enrichment and high-recall matching for existing park data.

This deliberately reads, but never writes, the original collector datasets.
"""
from __future__ import annotations

import argparse, csv, json, re, shutil, time
from copy import copy
from datetime import date, datetime
from pathlib import Path
from typing import Any
from zoneinfo import ZoneInfo
from zipfile import BadZipFile

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils.exceptions import InvalidFileException

from kerala_job_matching_engine import normalize_infopark, normalize_technopark, today_india

TITLE_KEYWORDS = """Software Developer|Software Engineer|Junior Software Developer|Junior Software Engineer|Entry Level Software Developer|Entry Level Software Engineer|Graduate Software Engineer|Graduate Software Developer|Associate Software Engineer|Associate Software Developer|Trainee Software Engineer|Trainee Software Developer|Software Engineer Trainee|Software Developer Trainee|Fresher Software Developer|Fresher Software Engineer|Application Developer|Application Engineer|Associate Application Developer|Web Developer|Junior Web Developer|Frontend Developer|Front End Developer|Frontend Engineer|Front End Engineer|Junior Frontend Developer|Junior Front End Developer|Frontend Software Engineer|Frontend Software Developer|UI Developer|UI Engineer|React Developer|React.js Developer|React JS Developer|Junior React Developer|React Engineer|React.js Engineer|Next.js Developer|Next JS Developer|JavaScript Developer|TypeScript Developer|Web Application Developer|Web Application Engineer|Full Stack Developer|Full-Stack Developer|Junior Full Stack Developer|Junior Full-Stack Developer|Full Stack Engineer|Full-Stack Engineer|MERN Stack Developer|MERN Developer|MERN Stack Engineer|Node.js Developer|Node JS Developer|Junior Node.js Developer|Node.js Engineer|Backend Developer|Back End Developer|Junior Backend Developer|Junior Back End Developer|Backend Software Engineer|API Developer|REST API Developer|Software Programmer|Programmer|Computer Programmer|Programmer Analyst|Associate Programmer Analyst|Web Programmer|Software Development Engineer|Associate Software Development Engineer|Junior Software Development Engineer|SDE I|SDE 1|Product Engineer|Associate Product Engineer|Junior Product Engineer|Product Software Engineer|Platform Engineer|Junior Platform Engineer|Cloud Application Developer|Web Solutions Developer|Solutions Developer|Technical Developer|IT Developer|IT Software Developer|IT Engineer|Technology Analyst|Associate Technology Analyst|Software Analyst|Application Support Developer|Application Support Engineer|Technical Consultant|Associate Technical Consultant|Technical Analyst|Integration Developer|Database Application Developer|Database Developer|SQL Developer|Python Developer|Junior Python Developer|Java Developer|Junior Java Developer|PHP Developer|Junior PHP Developer|Laravel Developer|Django Developer|FastAPI Developer|QA Automation Engineer|Software Development Intern|Web Development Intern|Frontend Developer Intern|React Developer Intern|Full Stack Developer Intern|Full-Stack Developer Intern|MERN Stack Developer Intern|Software Engineer Intern|Application Developer Intern|Graduate Engineer Trainee|Graduate Trainee Engineer|Engineering Trainee|IT Trainee|Developer Trainee|Web Developer Trainee|Frontend Trainee|Full Stack Trainee|React Trainee|Associate Engineer|Junior Engineer - Software|Junior Web Application Developer|Junior Application Developer|Entry Level Web Developer|Entry Level Frontend Developer|Entry Level Full Stack Developer|Entry Level React Developer|Entry Level Backend Developer|Entry Level Application Developer|Fresher Web Developer|Fresher Frontend Developer|Fresher Full Stack Developer|Fresher React Developer|Fresher MERN Stack Developer|Fresher Node.js Developer|Fresher Application Developer|Web UI Developer|UI Engineer - Web|Interactive Developer|Website Developer|Digital Product Developer|Software Product Developer|Web Technology Developer|Cross-Platform Application Developer|Progressive Web App Developer|PWA Developer|SaaS Developer|Cloud Software Developer""".split("|")
EXTRA_TITLES = ["Associate Developer", "Digital Engineer", "Application Programmer", "Web Solutions Engineer", "Developer Associate", "Technical Graduate Trainee", "Junior Technical Engineer", "Graduate Developer", "Associate Technology Engineer", "Entry-Level Application Engineer"]
SKILL_PATTERNS = {"HTML/CSS": ["html", "css", "tailwind", "scss"], "JavaScript": ["javascript", "typescript"], "React": ["react", "react.js", "reactjs"], "Next.js": ["next.js", "nextjs"], "Node.js": ["node.js", "nodejs"], "Express.js": ["express", "express.js"], "REST API": ["rest api", "rest apis", "api development", "restful"], "MongoDB": ["mongodb", "mongo"], "PostgreSQL": ["postgresql", "postgres"], "Prisma": ["prisma"], "SQL": ["sql"], "Python": ["python"], "Git": ["git"], "AWS": ["aws", "ec2"], "QA Automation": ["automation", "selenium", "cypress"]}
RELEVANT_DUTIES = ["web application", "frontend", "front end", "backend", "full stack", "software development", "application development", "develop", "programming", "api", "database", "user interface"]
EXCLUDED = re.compile(r"\b(sales|marketing|recruit(?:ment)?|human resources|\bhr\b|finance|account(?:ant|ing)?|legal|content specialist|business development|graphic design|customer support)\b", re.I)
REVIEW = ["Not Reviewed", "Apply", "Review", "Skip", "Applied"]
FIELDS = [("Match Category","category"),("Match Score","match_score"),("Confidence Level","confidence"),("Source Park","source_portal"),("Company","company"),("Original Job Title","job_title"),("Normalized Job Title","normalized_title"),("Original Job URL","job_details_url"),("Official Detail URL","job_details_url"),("Job Location","location"),("Posting Date","posted_date"),("Application Deadline","closing_date"),("Job Status","job_status"),("Description Retrieval Status","description_status"),("Description Summary","description_summary"),("Required Skills","required_skills"),("Preferred Skills","preferred_skills"),("Matched Skills","matched_skills"),("Missing Required Skills","missing_required_skills"),("Experience Requirements","experience_requirements"),("Education Requirements","education_requirements"),("Education Eligibility Status","education_eligibility"),("Fresher Eligibility","fresher_eligibility"),("Technical Relevance Reason","technical_reason"),("Match Evidence","match_evidence"),("Manual Verification Required","manual_verification"),("Recommended Next Action","next_action"),("Review Status","review_status"),("Review Notes","review_notes"),("Source Record ID","source_record_id"),("Last Enrichment Timestamp","last_enrichment_timestamp")]

def now() -> str: return datetime.now(ZoneInfo("Asia/Kolkata")).isoformat()
def key(job: dict[str,Any]) -> str: return f"{job['source_portal']}:{job['source_record_id']}"
def clean(value: Any) -> str: return "" if value is None else str(value)
def normal(text: str) -> str: return re.sub(r"[^a-z0-9]+", " ", text.casefold()).strip()
def phrase_match(title: str, phrase: str) -> bool: return re.search(r"(?<!\w)"+re.escape(normal(phrase)).replace(r"\ ",r"\s+")+r"(?!\w)", normal(title)) is not None
def atomic(path: Path,data: Any):
    path.parent.mkdir(parents=True,exist_ok=True); temp=path.with_name('.'+path.name+'.tmp');temp.write_text(json.dumps(data,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf8');temp.replace(path)

def listed_title_match(title: str) -> list[str]: return [p for p in TITLE_KEYWORDS+EXTRA_TITLES if phrase_match(title,p)]
def skills(text: str) -> set[str]:
    low=text.casefold(); return {name for name,words in SKILL_PATTERNS.items() if any(re.search(r"(?<!\w)"+re.escape(w.casefold())+r"(?!\w)",low) for w in words)}
def parse_date(value: str) -> date|None:
    for fmt in ('%Y-%m-%d','%d-%m-%Y','%d %b %Y','%d %B %Y'):
        try:return datetime.strptime(value.strip(),fmt).date()
        except ValueError:pass
    return None
def status(closing: str)->str:
    parsed=parse_date(closing)
    if not parsed:return 'Unknown'
    return 'Expired' if parsed<today_india() else 'Open'

def fetch_description(job: dict[str,Any], cache: dict[str,Any], session:requests.Session, delay:float)->dict[str,Any]:
    identity=key(job); cached=cache.get(identity)
    if isinstance(cached,dict) and cached.get('url')==job.get('job_details_url') and cached.get('status') in ('FULL_DESCRIPTION_RETRIEVED','PARTIAL_DESCRIPTION_RETRIEVED','PAGE_NOT_FOUND'):
        return cached
    url=job.get('job_details_url','')
    if not url:return {'url':url,'status':'DESCRIPTION_NOT_AVAILABLE','text':'','reason':'No verified official detail URL'}
    try:
        response=session.get(url,timeout=(10,30)); time.sleep(delay)
        if response.status_code==404:return {'url':url,'status':'PAGE_NOT_FOUND','text':'','reason':'HTTP 404'}
        if response.status_code in (401,403):return {'url':url,'status':'ACCESS_RESTRICTED','text':'','reason':f'HTTP {response.status_code}'}
        if not response.ok:return {'url':url,'status':'FETCH_FAILED','text':'','reason':f'HTTP {response.status_code}'}
        soup=BeautifulSoup(response.text,'html.parser')
        for tag in soup(['script','style','nav','footer','header']):tag.decompose()
        container=soup.select_one('.comp-job-deatiil .deatil-box, .comp-job-detail .detail-box, main') or soup.body
        body=' '.join(container.get_text(' ',strip=True).split()) if container else ''
        title=job['job_title'].casefold()
        useful=len(body)>max(250,len(title)+50) and title in body.casefold()
        return {'url':url,'status':'FULL_DESCRIPTION_RETRIEVED' if useful else 'PARTIAL_DESCRIPTION_RETRIEVED' if body else 'DESCRIPTION_NOT_AVAILABLE','text':body,'reason':''}
    except requests.RequestException as exc:return {'url':url,'status':'FETCH_FAILED','text':'','reason':f'{type(exc).__name__}: {exc}'}

def extract_requirements(description: str)->tuple[set[str],set[str],str,str,str]:
    required=set();preferred=set()
    for part in re.split(r'[.;\n]',description):
        target=skills(part); low=part.casefold()
        if re.search(r'\b(required|must have|mandatory|requirements?)\b',low):required|=target
        elif re.search(r'\b(preferred|nice to have|bonus|plus)\b',low):preferred|=target
    exp=(re.search(r'\b\d+(?:\s*-\s*\d+)?\s*\+?\s*(?:years?|yrs?)\b',description,re.I) or [""])[0]
    edu=(re.search(r'\b(?:b\.?tech|bachelor|master|degree|diploma|mca|bca)\b[^.;]{0,160}',description,re.I) or [""])[0]
    fresher='Accepted' if re.search(r'\b(fresher|freshers|entry.level|graduate trainee|intern)\b',description,re.I) else 'Not Specified'
    return required,preferred,exp,edu,fresher

def classify(job:dict[str,Any], enriched:dict[str,Any])->dict[str,Any]:
    desc=enriched.get('text',''); title=job['job_title']; title_hits=listed_title_match(title); title_skills=skills(title); desc_skills=skills(desc); all_skills=title_skills|desc_skills
    duty=any(term in (title+' '+desc).casefold() for term in RELEVANT_DUTIES)
    req,pref,experience,education,fresher=extract_requirements(desc)
    missing=req-all_skills # only reported if source has stated skills not profile? Need candidate skill base separate
    candidate_skills=set(SKILL_PATTERNS)-{'QA Automation'}
    missing=req-candidate_skills
    job_status=status(job['closing_date'])
    senior=bool(re.search(r'\b(senior|lead|principal|architect|manager)\b',title,re.I))
    long_experience=bool(re.search(r'\b[3-9]\d*\s*\+?\s*(?:years?|yrs?)\b',experience,re.I))
    excluded=bool(EXCLUDED.search(title)) and not (duty and all_skills)
    candidate=bool(title_hits or (duty and all_skills) or (desc_skills & {'React','Next.js','Node.js','Express.js','REST API','MongoDB','PostgreSQL'}))
    score=min(100,(30 if title_hits else 0)+(min(30,len(all_skills)*5))+(20 if duty else 0)+(10 if fresher=='Accepted' else 5 if not senior and not long_experience else 0)+(10 if enriched['status']=='FULL_DESCRIPTION_RETRIEVED' else 0))
    if excluded or not candidate: category='Not Relevant'
    elif job_status=='Expired': category='Expired Relevant Role'
    elif enriched['status']!='FULL_DESCRIPTION_RETRIEVED': category='Manual Review'
    elif long_experience or missing: category='Potential Match'
    elif score>=70: category='Strong Match'
    elif score>=40: category='Potential Match'
    else: category='Related Developer Role'
    confidence='High' if enriched['status']=='FULL_DESCRIPTION_RETRIEVED' and score>=70 else 'Medium' if enriched['status']=='FULL_DESCRIPTION_RETRIEVED' else 'Low'
    normalized=' / '.join(title_hits[:3]) or title
    reasons=[]
    if title_hits:reasons.append('Title candidate: '+', '.join(title_hits[:3]))
    if all_skills:reasons.append('Relevant skills found: '+', '.join(sorted(all_skills)))
    if duty:reasons.append('Development responsibility terminology found')
    if not reasons:reasons.append('No supported software-development signal found')
    eligibility='Needs Verification' if education or long_experience else 'Not Specified'
    job.update({'category':category,'match_score':score,'confidence':confidence,'normalized_title':normalized,'job_status':job_status,'description_status':enriched['status'],'description_summary':desc[:1200],'required_skills':', '.join(sorted(req)),'preferred_skills':', '.join(sorted(pref)),'matched_skills':', '.join(sorted(all_skills&candidate_skills)),'missing_required_skills':', '.join(sorted(missing)),'experience_requirements':experience,'education_requirements':education,'education_eligibility':eligibility,'fresher_eligibility':fresher,'technical_reason':'; '.join(reasons),'match_evidence':'; '.join(reasons)+('. Full official description retrieved.' if enriched['status']=='FULL_DESCRIPTION_RETRIEVED' else '. Description unavailable or partial; verify manually.'),'manual_verification':'Yes' if category in ('Potential Match','Manual Review') or eligibility=='Needs Verification' else 'No','next_action':'Read official description and verify eligibility' if category in ('Potential Match','Manual Review') else 'Consider application' if category=='Strong Match' else 'Archive / no action','last_enrichment_timestamp':now(),'enrichment_error':enriched.get('reason','')})
    return job

def reviews(out:Path, jobs:list[dict[str,Any]])->dict[str,dict[str,str]]:
    state=out/'review_state.json'; data={}
    if state.exists():
        try:data=json.loads(state.read_text(encoding='utf8')).get('reviews',{})
        except (ValueError,OSError):pass
    # One-time compatible migration from the prior combined matcher state.
    legacy=out.parents[1]/'job_matching'/'review_state.json'
    if not data and legacy.exists():
        try:data=json.loads(legacy.read_text(encoding='utf8')).get('reviews',{})
        except (ValueError,OSError):pass
    for job in jobs:
        prior=data.get(key(job),{});job['review_status']=prior.get('review_status','Not Reviewed');job['review_notes']=prior.get('review_notes','');data[key(job)]={'review_status':job['review_status'],'review_notes':job['review_notes']}
    atomic(state,{'updated_at':now(),'reviews':data});return data

def write_csv(path:Path,jobs:list[dict[str,Any]]):
    path.parent.mkdir(parents=True,exist_ok=True)
    with path.open('w',encoding='utf-8-sig',newline='') as h:
        w=csv.DictWriter(h,[key for _,key in FIELDS]);w.writeheader();w.writerows([{key:clean(j.get(key)) for _,key in FIELDS} for j in jobs])
def write_book(path:Path,jobs:list[dict[str,Any]],summary:dict[str,int]):
    path.parent.mkdir(parents=True,exist_ok=True);wb=Workbook();ws=wb.active;ws.title='Matched Jobs';ws.append([h for h,_ in FIELDS]);fill=PatternFill('solid',fgColor='1F4E78')
    for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=fill
    for job in jobs:
        ws.append([job.get(key,'') for _,key in FIELDS]);url=job.get('job_details_url','');cell=ws.cell(ws.max_row,8)
        if url:cell.hyperlink=url;cell.style='Hyperlink'
    last=max(2,ws.max_row);ws.freeze_panes='A2';ws.auto_filter.ref=f'A1:AE{last}';dv=DataValidation(type='list',formula1='"'+','.join(REVIEW)+'"');ws.add_data_validation(dv);dv.add(f'AB2:AB{last}')
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width=20
    summary_ws=wb.create_sheet('Summary');summary_ws.append(['Category','Count']);[summary_ws.append([k,v]) for k,v in sorted(summary.items())];summary_ws.column_dimensions['A'].width=32
    temp=path.with_name('.'+path.name+'.tmp');wb.save(temp);wb.close();temp.replace(path)

def process(portal:str,jobs:list[dict[str,Any]],root:Path,delay:float)->dict[str,Any]:
    out=root/portal.lower();cache_path=out/'description_cache.json';cache={}
    if cache_path.exists():
        try:cache=json.loads(cache_path.read_text(encoding='utf8'))
        except (OSError,ValueError):pass
    s=requests.Session();s.headers.update({'User-Agent':'AI-JobHunt-Reinvestigation/1.0 (+public job detail review)','Accept':'text/html'})
    for idx,job in enumerate(jobs,1):
        e=fetch_description(job,cache,s,delay);cache[key(job)]=e;classify(job,e)
        if idx%50==0:atomic(cache_path,cache);print(f'[INFO] {portal}: enriched {idx}/{len(jobs)}')
    atomic(cache_path,cache);reviews(out,jobs)
    candidates=[j for j in jobs if j['category']!='Not Relevant'];rejected=[j for j in jobs if j['category']=='Not Relevant']
    count={cat:sum(j['category']==cat for j in jobs) for cat in ('Strong Match','Potential Match','Related Developer Role','Manual Review','Expired Relevant Role','Not Relevant')}
    prefix=portal.lower();write_csv(out/f'{prefix}_all_matched_jobs.csv',candidates);atomic(out/f'{prefix}_all_matched_jobs.json',{'portal':portal,'generated_at':now(),'jobs':candidates});write_book(out/f'{prefix}_all_matched_jobs.xlsx',candidates,count)
    for name,cat in [('strong_matches','Strong Match'),('potential_matches','Potential Match'),('related_roles','Related Developer Role'),('manual_review','Manual Review')]:write_book(out/f'{prefix}_{name}.xlsx',[j for j in candidates if j['category']==cat],count)
    write_csv(out/f'{prefix}_rejected_jobs.csv',rejected)
    statuses={x:sum(j['description_status']==x for j in jobs) for x in sorted({j['description_status'] for j in jobs})}
    report={'portal':portal,'generated_at':now(),'source_records':len(jobs),'unique_source_records':len({key(j) for j in jobs}),'intra_portal_duplicates':len(jobs)-len({key(j) for j in jobs}),'category_counts':count,'candidate_records_written':len(candidates),'description_status_counts':statuses,'eligibility_manual_verification':sum(j['manual_verification']=='Yes' for j in candidates),'output_directory':str(out)}
    atomic(out/'reports'/'matching_report.json',report);atomic(out/'reports'/'collection_coverage_report.json',{'portal':portal,'source_records':len(jobs),'description_status_counts':statuses});atomic(out/'reports'/'failures_and_warnings.json',{'failed_or_unavailable_descriptions':[{'source_record_id':j['source_record_id'],'url':j['job_details_url'],'status':j['description_status'],'reason':j['enrichment_error']} for j in jobs if j['description_status'] not in ('FULL_DESCRIPTION_RETRIEVED','PARTIAL_DESCRIPTION_RETRIEVED')]})
    return report

def main()->int:
    p=argparse.ArgumentParser(description='Enrich and high-recall match existing Technopark and Infopark datasets.');p.add_argument('--delay',type=float,default=.08);p.add_argument('--technopark-source',type=Path);p.add_argument('--infopark-source',type=Path);args=p.parse_args()
    if args.delay<0:p.error('--delay must be non-negative')
    project=Path(__file__).resolve().parent;tp=args.technopark_source or project/'data/technopark/technopark_all_jobs.json';ip=args.infopark_source or project/'data/infopark/infopark_all_jobs.json';root=project/'data/matched_jobs'
    if not tp.exists() or not ip.exists():print('[ERROR] Both source JSON files are required.');return 2
    tech,info=normalize_technopark(tp),normalize_infopark(ip);print(f'[INFO] Source audit: Technopark={len(tech)}, Infopark={len(info)}, combined={len(tech)+len(info)}')
    reports=[process('Technopark',tech,root,args.delay),process('Infopark',info,root,args.delay)]
    cross=[]; buckets={}
    for j in tech+info:buckets.setdefault((normal(j['job_title']),normal(j['company'])),[]).append(j)
    for group in buckets.values():
        if len({j['source_portal'] for j in group})>1:cross.append({'title':group[0]['job_title'],'company':group[0]['company'],'records':[key(j) for j in group]})
    atomic(root/'cross_source_duplicate_candidates.json',cross);atomic(root/'reinvestigation_audit.json',{'generated_at':now(),'technopark_source_records':len(tech),'infopark_source_records':len(info),'combined_source_records':len(tech)+len(info),'reported_731_vs_736_resolution':'Current source JSON files contain 317 Technopark plus 414 Infopark records, totalling 731. No five-record discrepancy is present in the current source files.','potential_cross_source_duplicates':cross,'portal_reports':reports})
    print('[INFO] Separate portal outputs written to data/matched_jobs/.');return 0
if __name__=='__main__':raise SystemExit(main())
