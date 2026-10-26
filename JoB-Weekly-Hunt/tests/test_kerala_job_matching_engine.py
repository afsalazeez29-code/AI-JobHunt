import json
from datetime import date

from openpyxl import load_workbook

from kerala_job_matching_engine import (
    counts, deadline, evaluate, normalize_infopark, normalize_technopark,
    potential_duplicates, read_review_data, requirement_sections, role_family,
    skill_mentions, write_workbook,
)


def source_job(title="React Developer", description="", portal="Technopark"):
    return {
        "source_portal": portal, "source_record_id": "1", "job_listing_id": "JOB-1", "job_title": title,
        "company": "Example", "posted_date": "2026-10-24", "closing_date": "2026-11-10", "job_description": description,
        "required_skills": "", "preferred_skills": "", "experience_requirement": "", "education_requirement": "",
        "employment_type": "", "walk_in_status": "", "walk_in_date_time": "", "location": "", "work_mode": "", "salary": "",
        "job_details_url": "https://technopark.in/job-details/1", "original_source_file": "fixture", "original_source_record": {"id": 1},
    }


def test_normalizes_real_portal_shapes(tmp_path):
    tp_path = tmp_path / "tp.json"; ip_path = tmp_path / "ip.json"
    tp_path.write_text(json.dumps({"collection_metadata": {"job_details_url_validation": {"validated": True}}, "jobs": [{"id": 7, "job_title": "React Developer", "company": {"company": "Co"}}]}), encoding="utf-8")
    ip_path.write_text(json.dumps([{ "job_id": "9", "job_title": "Node.js Developer", "company_name": "IP Co", "date_posted": "25-10-2026", "application_deadline": "28 Oct 2026", "job_details_url": "https://infopark.in/company-jobs/details/1/9" }]), encoding="utf-8")
    tp, ip = normalize_technopark(tp_path), normalize_infopark(ip_path)
    assert tp[0]["source_portal"] == "Technopark" and tp[0]["job_details_url"].endswith("/7")
    assert ip[0]["source_portal"] == "Infopark" and ip[0]["source_record_id"] == "9"


def test_aliases_and_required_vs_preferred_skills():
    assert {"React", "Next.js", "Node.js"}.issubset(skill_mentions("ReactJS, NextJS and Node.js"))
    required, preferred = requirement_sections("Required: React and TypeScript. Nice to have: Prisma.")
    assert required == {"React", "TypeScript"} and preferred == {"Prisma"}


def test_title_only_relevant_role_is_manual_review_with_low_coverage():
    result = evaluate(source_job(), date(2026, 10, 26))
    assert result["classification"] == "Manual Review"
    assert result["matching_method"] == "Title-Based" and result["evidence_coverage"] == "20%"
    assert "description" in result["missing_information"].lower()


def test_description_based_conflict_is_not_overridden_by_score():
    result = evaluate(source_job("Full-Stack Developer", "Required: React, Node.js. Must have 3+ years experience."), date(2026, 10, 26))
    assert result["classification"] == "Low Match / Excluded"
    assert "Professional experience requirement" in result["mandatory_conflicts"]


def test_deadline_and_seniority_classification():
    assert deadline("2026-10-25", date(2026, 10, 26)) == ("Expired", -1)
    assert deadline("28 Oct 2026", date(2026, 10, 26))[0] == "Closing Soon"
    assert role_family("Senior React Developer")[1] == 25


def test_cross_source_duplicates_are_flagged_not_merged():
    a, b = source_job(portal="Technopark"), source_job(portal="Infopark")
    b["source_record_id"] = "2"
    duplicates = potential_duplicates([a, b])
    assert len(duplicates) == 1 and "Technopark:1" in duplicates[0]["source_keys"] and "Infopark:2" in duplicates[0]["source_keys"]


def test_workbook_preserves_status_and_notes_and_has_required_sheets(tmp_path):
    job = evaluate(source_job(), date(2026, 10, 26)); job.update({"review_status": "Not Reviewed", "notes": ""})
    report = {"execution_status": "COMPLETE", "completed_at": "2026-10-26T00:00:00+05:30", "counts": counts([job])}
    workbook = tmp_path / "match.xlsx"; write_workbook(workbook, [job], report)
    wb = load_workbook(workbook); ws = wb["All Shortlisted Jobs"]; ws["A2"] = "Apply"; ws["AB2"] = "follow up"; wb.save(workbook); wb.close()
    reviews, warnings = read_review_data(workbook, tmp_path / "state.json", tmp_path / "backups", "test")
    assert not warnings and reviews["Technopark:1"] == {"review_status": "Apply", "notes": "follow up"}
    wb = load_workbook(workbook); assert {"Strong Matches", "Related Matches", "Manual Review", "Expired - Past Deadline", "All Shortlisted Jobs", "Summary"}.issubset(wb.sheetnames)
    assert wb["All Shortlisted Jobs"].freeze_panes == "A2" and wb["All Shortlisted Jobs"].data_validations.dataValidation; wb.close()
