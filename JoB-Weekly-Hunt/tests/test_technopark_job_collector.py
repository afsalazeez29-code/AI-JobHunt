import csv
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

from technopark_job_collector import (
    API_ENDPOINT,
    CollectionResult,
    FetchResult,
    TechnoparkCollector,
    compute_status,
    deduplicate_records,
    is_valid_next_url,
    make_normalized,
    read_existing_reviews,
    validate_outputs,
    write_csv,
    write_json,
    write_workbook,
)


def job(job_id, title="Engineer"):
    return {
        "id": job_id,
        "job_listing_id": f"JOB-{job_id}",
        "job_title": title,
        "posted_date": "2026-10-24",
        "is_walk_in": 0,
        "closing_date": "2026-11-10",
        "company": {"id": 99, "company": "Example, Ltd", "logo": "/logo.png"},
    }


class Response:
    def __init__(self, payload, status_code=200):
        self.payload = payload
        self.status_code = status_code
        self.ok = status_code < 400
        self.reason = "test"
        self.headers = {"Content-Type": "application/json"}

    def json(self):
        if self.payload == "bad-json":
            raise ValueError("bad JSON")
        return self.payload


class Session:
    def __init__(self, pages):
        self.pages = pages
        self.headers = {}

    def get(self, url, **_kwargs):
        page = int(url.split("page=")[1])
        return Response(self.pages[page])


def test_only_expected_next_urls_are_accepted():
    assert is_valid_next_url(f"{API_ENDPOINT}?page=2")
    assert not is_valid_next_url("http://technopark.in/api/paginated-jobs?page=2")
    assert not is_valid_next_url("https://example.test/api/paginated-jobs?page=2")
    assert not is_valid_next_url("https://technopark.in/api/paginated-jobs?page=0")
    assert not is_valid_next_url("https://technopark.in/something?page=2")


def test_collects_all_pages_and_uses_last_page_fallback(tmp_path):
    pages = {
        1: {"current_page": 1, "last_page": 2, "next_page_url": None, "total": 2, "per_page": 1, "data": [job(1)]},
        2: {"current_page": 2, "last_page": 2, "next_page_url": None, "total": 2, "per_page": 1, "data": [job(2)]},
    }
    collector = TechnoparkCollector(tmp_path, delay=0)
    collector.session = Session(pages)
    result = collector.collect("test")
    assert [page.page_number for page in result.page_results] == [1, 2]
    assert [record["id"] for record in result.raw_records] == [1, 2]
    assert (tmp_path / "raw" / "page_01.json").exists()


def test_deduplicate_ids_but_retain_missing_identifiers():
    first = job(1, "Original")
    duplicate = job(1, "Changed")
    missing = {"job_title": "No identifiers", "company": {"company": "Example"}}
    unique, report = deduplicate_records([first, duplicate, missing])
    assert unique == [first, missing]
    assert report["duplicate_occurrences"] == 1
    assert report["duplicate_conflicting_identifiers"] == ["id:1"]
    assert len(report["missing_stable_identifier_records"]) == 1


def test_failed_page_makes_collection_partial():
    first = {"current_page": 1, "last_page": 2, "next_page_url": f"{API_ENDPOINT}?page=2", "total": 2, "per_page": 1, "data": [job(1)]}
    result = CollectionResult(
        raw_records=[job(1)],
        page_results=[FetchResult(1, f"{API_ENDPOINT}?page=1", True, first), FetchResult(2, f"{API_ENDPOINT}?page=2", False, error="HTTP 503")],
        first_payload=first,
        observed_pagination=[{"current_page": 1, "last_page": 2, "next_page_url": f"{API_ENDPOINT}?page=2", "total": 2, "per_page": 1}],
        warnings=[],
    )
    status, reasons = compute_status(result, 1)
    assert status == "PARTIAL"
    assert reasons


def test_writes_and_reopens_all_formats_and_preserves_reviews(tmp_path):
    records = [job(1, 'A "quoted", Unicode café'), job(2, "Second")]
    normalized = make_normalized(records, {"id:1": "https://technopark.in/job-details/1", "id:2": "https://technopark.in/job-details/2"})
    json_path = tmp_path / "technopark_all_jobs.json"
    csv_path = tmp_path / "technopark_all_jobs.csv"
    workbook_path = tmp_path / "technopark_jobs_review.xlsx"
    write_json(json_path, {"collection_status": "COMPLETE"}, records)
    write_csv(csv_path, normalized)
    write_workbook(workbook_path, normalized, records, {})
    workbook = load_workbook(workbook_path)
    worksheet = workbook["Technopark Jobs"]
    worksheet["A2"] = "Apply"
    worksheet["Q2"] = "good fit"
    workbook.save(workbook_path)
    workbook.close()
    previous, warnings = read_existing_reviews(workbook_path, tmp_path / "raw" / "workbook_backups", "test")
    assert not warnings
    assert previous["id:1"] == ("Apply", "good fit")
    write_workbook(workbook_path, normalized, records, previous)
    result = validate_outputs(json_path, csv_path, workbook_path, 2)
    assert result["passed"], result["errors"]
    workbook = load_workbook(workbook_path)
    worksheet = workbook["Technopark Jobs"]
    assert worksheet["A2"].value == "Apply"
    assert worksheet["Q2"].value == "good fit"
    assert worksheet["N2"].hyperlink.target == "https://technopark.in/job-details/1"
    assert worksheet.freeze_panes == "A2"
    assert worksheet.auto_filter.ref
    assert worksheet.data_validations.dataValidation
    workbook.close()
    with csv_path.open(encoding="utf-8-sig", newline="") as handle:
        rows = list(csv.DictReader(handle))
    assert rows[0]["job_title"] == 'A "quoted", Unicode café'
    assert json.loads(json_path.read_text(encoding="utf-8"))["jobs"][0]["company"]["company"] == "Example, Ltd"


def test_malformed_json_is_not_treated_as_a_success(tmp_path):
    collector = TechnoparkCollector(tmp_path, delay=0, max_retries=0)
    collector.session = type("BadSession", (), {"get": lambda *_args, **_kwargs: Response("bad-json")})()
    result = collector.fetch_page(f"{API_ENDPOINT}?page=1", 1)
    assert not result.success
    assert "Invalid JSON" in result.error
